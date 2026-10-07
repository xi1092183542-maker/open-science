"""Small, independently authored research document exporter (MIT)."""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import tempfile
from xml.sax.saxutils import escape


def validate(data):
    if not isinstance(data, dict) or not isinstance(data.get("title"), str):
        raise ValueError("title must be a string")
    sections = data.get("sections", [])
    sources = data.get("sources", [])
    if not isinstance(sections, list) or not isinstance(sources, list):
        raise ValueError("sections and sources must be lists")
    for section in sections:
        if not isinstance(section, dict) or not isinstance(section.get("heading"), str):
            raise ValueError("each section needs a heading")
        if not isinstance(section.get("paragraphs"), list) or not all(
            isinstance(p, str) for p in section["paragraphs"]
        ):
            raise ValueError("paragraphs must be strings")
    if not all(isinstance(s, str) for s in sources):
        raise ValueError("sources must be strings")
    table = data.get("table")
    if table is not None:
        if not isinstance(table, dict):
            raise ValueError("table must be an object")
        cols, rows = table.get("columns"), table.get("rows")
        if not isinstance(cols, list) or not cols or not all(isinstance(c, str) for c in cols):
            raise ValueError("table.columns must be nonempty text headers")
        if not isinstance(rows, list):
            raise ValueError("table.rows must be a list")
        for row in rows:
            if not isinstance(row, list) or len(row) != len(cols):
                raise ValueError("each row must match the header width")
            for value in row:
                if value is not None and type(value) not in (str, int, float):
                    raise ValueError("table values must be text, numbers or null")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("non-finite numbers are not valid measurements")
    return data


def all_sections(data):
    yield from data.get("sections", [])
    if data.get("sources"):
        yield {"heading": "Sources", "paragraphs": data["sources"]}


def text(value):
    return "" if value is None else str(value)


def docx(data, output):
    from docx import Document
    document = Document()
    document.add_heading(data["title"], 0)
    for section in all_sections(data):
        document.add_heading(section["heading"], 1)
        for paragraph in section["paragraphs"]:
            document.add_paragraph(paragraph)
    if table := data.get("table"):
        grid = document.add_table(rows=1, cols=len(table["columns"]))
        grid.style = "Table Grid"
        for cell, value in zip(grid.rows[0].cells, table["columns"]):
            cell.text = value
        for row in table["rows"]:
            for cell, value in zip(grid.add_row().cells, row):
                cell.text = text(value)
    document.save(output)


def xlsx(data, output):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    def append(target, values):
        target.append(values)
        for cell in target[target.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Untrusted text must not become a formula.

    append(sheet, [data["title"]])
    for section in all_sections(data):
        append(sheet, [section["heading"]])
        for paragraph in section["paragraphs"]:
            append(sheet, [paragraph])
    if table := data.get("table"):
        grid = workbook.create_sheet("Data")
        append(grid, table["columns"])
        for row in table["rows"]:
            append(grid, row)
        grid.freeze_panes = "A2"
        grid.auto_filter.ref = grid.dimensions
    workbook.save(output)


def pptx(data, output):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def slide(title, paragraphs):
        page = deck.slides.add_slide(deck.slide_layouts[1])
        page.shapes.title.text = title
        frame = page.placeholders[1].text_frame
        frame.clear()
        for i, paragraph in enumerate(paragraphs):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = paragraph
            p.font.size = Pt(20)
            p.font.name = "Microsoft YaHei"

    slide(data["title"], [])
    for section in all_sections(data):
        paragraphs = section["paragraphs"]
        if any(len(p) > 180 for p in paragraphs):
            raise ValueError("split PPTX paragraphs longer than 180 characters before export")
        for start in range(0, max(1, len(paragraphs)), 4):
            slide(section["heading"], paragraphs[start:start + 4])
    if table := data.get("table"):
        if len(table["columns"]) > 6:
            raise ValueError("PPTX tables support at most six columns; split wider data")
        for start in range(0, max(1, len(table["rows"])), 8):
            page = deck.slides.add_slide(deck.slide_layouts[5])
            page.shapes.title.text = "Data"
            values = [table["columns"], *table["rows"][start:start + 8]]
            grid = page.shapes.add_table(len(values), len(table["columns"]),
                                        Inches(.5), Inches(1.5), Inches(12), Inches(5)).table
            for i, row in enumerate(values):
                for j, value in enumerate(row):
                    if len(text(value)) > 70:
                        raise ValueError("PPTX table text exceeds layout limit; split the table")
                    cell = grid.cell(i, j)
                    cell.text = text(value)
                    for p in cell.text_frame.paragraphs:
                        p.font.size = Pt(14)
                        p.font.name = "Microsoft YaHei"
    deck.save(output)


def pdf(data, output):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "STSong-Light"
        style.wordWrap = "CJK"
    para = lambda value, style="BodyText": Paragraph(escape(text(value)), styles[style])
    flow = [para(data["title"], "Title")]
    for section in all_sections(data):
        flow.append(para(section["heading"], "Heading2"))
        flow.extend(para(p) for p in section["paragraphs"])
        flow.append(Spacer(1, 8))
    if table := data.get("table"):
        width = (595.28 - 96) / len(table["columns"])
        grid = LongTable([[para(v) for v in row] for row in
                          [table["columns"], *table["rows"]]],
                         colWidths=[width] * len(table["columns"]), repeatRows=1)
        grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        flow.append(grid)
    SimpleDocTemplate(str(output), leftMargin=48, rightMargin=48).build(flow)


def export(data, output: Path, overwrite=False):
    validate(data)
    writers = {".docx": docx, ".xlsx": xlsx, ".pptx": pptx, ".pdf": pdf}
    writer = writers.get(output.suffix.lower())
    if writer is None:
        raise ValueError("supported extensions: .docx, .xlsx, .pptx, .pdf")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Reserve destination exclusively so a collision never replaces an original.
    reserved = False
    if not overwrite:
        with output.open("xb"):
            pass
        reserved = True
    temporary = None
    try:
        fd, name = tempfile.mkstemp(dir=output.parent, suffix=output.suffix)
        os.close(fd)
        temporary = Path(name)
        writer(data, temporary)
        os.replace(temporary, output)
    except BaseException:
        if reserved:
            output.unlink(missing_ok=True)
        raise
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error("input and output must differ")
    export(json.loads(args.input.read_text(encoding="utf-8")), args.output, args.overwrite)
    print(args.output.resolve())
