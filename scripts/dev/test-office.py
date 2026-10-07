"""Round-trip document content and failure preservation checks; no Office required."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "runtime/skills/office/research-documents/scripts/export_document.py"
spec = importlib.util.spec_from_file_location("exporter", SCRIPT)
exporter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(exporter)

DATA = {
    "title": "Formulation study / 配方研究",
    "sections": [{"heading": "Observations", "paragraphs": ["Synthetic test data; not an experimental result."]}],
    "table": {"columns": ["Sample", "Mass (g)", "Note"],
              "rows": [["A", 12.5, "=1+1"], ["B", None, "<unmeasured>"]]},
    "sources": ["fixture:synthetic-001"],
}


class ExportTests(unittest.TestCase):
    def test_round_trip_all_formats(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory(prefix="chem-office-") as folder:
            paths = {ext: Path(folder) / ("report." + ext) for ext in ("docx", "xlsx", "pptx", "pdf")}
            for path in paths.values():
                exporter.export(DATA, path)
            document = Document(paths["docx"])
            self.assertIn(DATA["title"], [p.text for p in document.paragraphs])
            self.assertEqual(document.tables[0].cell(1, 1).text, "12.5")
            self.assertIn(DATA["sources"][0], [p.text for p in document.paragraphs])
            book = load_workbook(paths["xlsx"])
            self.assertEqual(book["Report"]["A1"].value, DATA["title"])
            self.assertEqual(book["Data"]["B2"].value, 12.5)
            self.assertEqual(book["Data"]["C2"].value, "=1+1")
            self.assertEqual(book["Data"]["C2"].data_type, "s")
            self.assertIsNone(book["Data"]["B3"].value)
            book.close()
            deck = Presentation(paths["pptx"])
            texts = [s.text for page in deck.slides for s in page.shapes if s.has_text_frame]
            self.assertIn(DATA["title"], texts)
            self.assertIn(DATA["sources"][0], texts)
            tables = [s.table for page in deck.slides for s in page.shapes if s.has_table]
            self.assertEqual(tables[0].cell(1, 1).text, "12.5")
            content = "\n".join(p.extract_text() for p in PdfReader(paths["pdf"]).pages)
            for expected in ("配方研究", "12.5", DATA["sources"][0]):
                self.assertIn(expected, content)

    def test_collision_and_failed_overwrite_preserve_original(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "existing.docx"
            path.write_bytes(b"original")
            with self.assertRaises(FileExistsError):
                exporter.export(DATA, path)
            with patch.object(exporter, "docx", side_effect=RuntimeError("simulated write failure")):
                with self.assertRaises(RuntimeError):
                    exporter.export(DATA, path, overwrite=True)
            self.assertEqual(path.read_bytes(), b"original")
            self.assertEqual(list(Path(folder).iterdir()), [path])

    def test_invalid_measurements_do_not_create_artifacts(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.xlsx"
            for value in (float("nan"), float("inf"), True, {"value": 2}):
                with self.subTest(value=value), self.assertRaises(ValueError):
                    exporter.export({"title": "Invalid", "table": {"columns": ["x"], "rows": [[value]]}}, path)
                self.assertFalse(path.exists())

    def test_pptx_layout_failure_cleans_new_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.pptx"
            with self.assertRaises(ValueError):
                exporter.export({"title": "Too long", "sections": [{"heading": "Text", "paragraphs": ["x" * 181]}]}, path)
            self.assertEqual(list(Path(folder).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
